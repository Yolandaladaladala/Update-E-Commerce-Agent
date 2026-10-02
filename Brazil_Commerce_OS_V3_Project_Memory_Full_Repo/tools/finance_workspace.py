from io import BytesIO
import json
from pathlib import Path
import pandas as pd
import streamlit as st
from tools.finance import analyse_finance, _normalise_columns, load_finance_file
from llm import chat

def load_book(file):
    if file.name.lower().endswith(('.xlsx','.xls')):
        sheets=pd.read_excel(file,sheet_name=None)
        return {n:_normalise_columns(d).dropna(how='all') for n,d in sheets.items() if n!='填写说明'}
    return {'订单':load_finance_file(file)}

def export_excel(results, sheets, advisor):
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    bio=BytesIO()
    with pd.ExcelWriter(bio,engine='openpyxl') as writer:
        summary=[];quality=[];details=[]
        for currency,c in results.items():
            for key in ['gross_sales','net_sales','order_count','aov','gross_profit','contribution_profit']:
                summary.append({'币种':currency,'指标':key,'金额或数量':c.get(key),'说明':'不完整' if c.get(key) is None else ('只扣已知扣减' if key=='net_sales' and not c['net_sales_complete'] else '')})
            quality += [dict(币种=currency,**r) for r in c['quality']]
            details.append(c['detail'])
            for dim,rows in c['group_analysis'].items():
                pd.DataFrame(rows).to_excel(writer,sheet_name=f'{currency}_{dim}'[:31],index=False)
        pd.DataFrame(summary).to_excel(writer,sheet_name='经营结果',index=False)
        pd.DataFrame(quality).to_excel(writer,sheet_name='数据缺口',index=False)
        pd.concat(details,ignore_index=True).to_excel(writer,sheet_name='计算明细',index=False)
        pd.DataFrame({'分析':advisor.splitlines()}).to_excel(writer,sheet_name='经营分析',index=False)
        for i,(n,d) in enumerate(sheets.items()): d.to_excel(writer,sheet_name=f'原始_{i+1}_{n}'[:31],index=False)
        wb=writer.book
        wb._sheets.sort(key=lambda s: 0 if s.title=='经营结果' else 1)
        for ws in wb:
            ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
            for cell in ws[1]: cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='17365D');cell.alignment=Alignment(wrap_text=True,vertical='center')
            ws.row_dimensions[1].height=32
            for cells in ws.columns:
                width=min(60,max(16,max(len(str(c.value or '')) for c in cells[:30])+2))
                ws.column_dimensions[get_column_letter(cells[0].column)].width=width
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment=Alignment(vertical='top',wrap_text=True)
                    if isinstance(cell.value,(int,float)):cell.number_format='#,##0.00'
    return bio.getvalue()

def render_finance(base, language_instruction, word_builder, pdf_builder, save_history):
    st.header('财务与经营分析 / Finance & Performance')
    st.caption('先算已知数据，再列出缺口；不同币种分别计算。此版本输出管理分析，不凭订单生成三大表。')
    template=Path(base)/'data'/'seller_finance_input_template.xlsx'
    if template.exists():st.download_button('下载卖家填写模板',template.read_bytes(),'seller_finance_input_template.xlsx')
    file=st.file_uploader('上传订单CSV或多Sheet Excel',type=['csv','xlsx','xls','json'],key='finance_v4_file')
    if not file:return
    try:sheets=load_book(file)
    except Exception as e:st.error(f'读取失败：{e}');return
    names=[n for n,d in sheets.items() if not d.empty]
    if not names:st.info('模板尚未填写，请先填写订单。');return
    order_name=st.selectbox('选择订单表',names,index=names.index('订单') if '订单' in names else 0)
    orders=sheets[order_name].copy()
    st.dataframe(orders.head(20),hide_index=True)
    a,b=st.columns(2)
    price_basis=a.selectbox('price字段代表什么？',['订单/该行总额','单件售价'])
    cost_basis=b.selectbox('cost/product_cost字段代表什么？',['该行总成本','单件成本'])
    currency=st.text_input('文件没有币种列时，填写币种（如BRL）','BRL')
    zero=st.multiselect('仅选择你确认整份数据没有发生的项目；未知不要选', ['discounts','refunds','platform_fee','logistics_cost','ad_spend','creator_cost','tax'])
    extra=[n for n in names if n!=order_name]
    fees=st.selectbox('订单费用表（可选）',['不合并']+extra)
    if st.button('计算并生成报告',type='primary'):
        try:
            if fees!='不合并':
                f=sheets[fees].copy()
                if 'order_no' not in f or 'order_no' not in orders:raise ValueError('合并费用需要两张表都有order_no。')
                if f['order_no'].duplicated().any() or orders['order_no'].duplicated().any():raise ValueError('订单费用自动合并仅支持一对一；发现重复订单，请先确认商品行和费用分摊。')
                cols=[c for c in f if c not in orders or c=='order_no']
                orders=orders.merge(f[cols],on='order_no',how='left',validate='one_to_one')
            if 'currency' not in orders:orders['currency']=currency.strip().upper()
            if orders['currency'].isna().any() or orders['currency'].astype(str).str.strip().eq('').any():raise ValueError('币种有空白，请补齐。')
            results={}
            for cur,g in orders.groupby('currency'):
                results[str(cur)]=analyse_finance(g, 'unit' if price_basis=='单件售价' else 'total','unit' if cost_basis=='单件成本' else 'total',zero)
            facts={cur:{k:v for k,v in c.items() if k not in ['detail']} for cur,c in results.items()}
            fallback='\n'.join(f'{cur}: 净销售额（已知扣减口径）{c["net_sales"]:,.2f}；完整贡献利润：'+('未能确认' if c['contribution_profit'] is None else f'{c["contribution_profit"]:,.2f}')+'。优先补充：'+', '.join(c['missing_for_profitability']) for cur,c in results.items())
            try:
                advisor=chat(language_instruction+'\nExplain the supplied deterministic financial results. Compare supported platform/SKU results, identify evidence-backed drivers and 3 prioritised actions. Do not invent numbers, attributed ROAS, taxes or missing costs. No cross-currency totals. State that period expenses and cash/balance sheets below were retained but not included in calculations. Facts:\n'+json.dumps(facts,ensure_ascii=False,default=str),max_tokens=2400)
            except Exception:advisor=fallback+'\nAI解释暂不可用；计算结果仍可下载。'
            if extra:advisor+='\n\n期间费用、收付款与余额等额外Sheet仅保留在原始数据中；本版尚未纳入经营利润或三大表计算。'
            st.session_state.finance_v4={'results':results,'advisor':advisor,'sheets':sheets,'file':file.name}
            save_history('Finance',file.name,'Completeness-aware analysis')
        except Exception as e:st.error(str(e))
    result=st.session_state.get('finance_v4')
    if not result:return
    if result['file']!=file.name:return
    for cur,c in result['results'].items():
        st.subheader(cur)
        a,b,d=st.columns(3)
        a.metric('销售额（已知扣减后）',f'{c["net_sales"]:,.2f}')
        b.metric('完整毛利','缺数据' if c['gross_profit'] is None else f'{c["gross_profit"]:,.2f}')
        d.metric('完整贡献利润','缺数据' if c['contribution_profit'] is None else f'{c["contribution_profit"]:,.2f}')
        st.dataframe(pd.DataFrame(c['quality']),hide_index=True)
        for dim,rows in c['group_analysis'].items():
            st.markdown(f'**{dim} 比较**');st.dataframe(pd.DataFrame(rows),hide_index=True)
    st.markdown(result['advisor'])
    st.download_button('下载计算好的Excel',export_excel(result['results'],result['sheets'],result['advisor']),'finance_analysis.xlsx')
    text=result['advisor']+'\n\n'+ '\n'.join(f'{cur}: '+json.dumps({k:c[k] for k in ['net_sales','gross_profit','contribution_profit','missing_for_profitability']},ensure_ascii=False) for cur,c in result['results'].items())
    try:
        st.download_button('下载Word分析',word_builder('财务经营分析',text),'finance_analysis.docx')
        st.download_button('下载PDF分析',pdf_builder('财务经营分析',text),'finance_analysis.pdf')
    except Exception as e:st.warning(f'文档导出失败，Excel仍可下载：{e}')
