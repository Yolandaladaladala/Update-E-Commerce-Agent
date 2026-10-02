"""Native Word/PDF research reports, sharing a small Markdown block parser."""
from __future__ import annotations
from io import BytesIO
import re
from html import escape
from datetime import datetime
from pathlib import Path
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib import font_manager
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image as RLImage, PageBreak, Table, TableStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

FONT_PATH=Path(__file__).resolve().parents[1] / 'data' / 'fonts' / 'NotoSansSC-Regular.ttf'

LABELS={
 'zh':{'subtitle':'市场进入评估报告','question':'研究问题','references':'参考来源','date':'研究日期','note':'证据检查说明'},
 'en':{'subtitle':'Market entry assessment','question':'Research question','references':'References','date':'Research date','note':'Evidence checks'},
 'pt-BR':{'subtitle':'Avaliação de entrada no mercado','question':'Pergunta de pesquisa','references':'Referências','date':'Data da pesquisa','note':'Verificação de evidências'}}

def _labels(language): return LABELS.get(language,LABELS['en'])

def _source_id(src,index): return src.get('source_id') or f'S{index}'

def _public_objective(result):
    return re.split(r'\n\s*(?:OUTPUT LANGUAGE REQUIREMENT|META INSTRUCTION)',result.get('objective',''))[0].strip()

def _clean_text(text):
    text=re.sub(r'<!--.*?-->','',str(text),flags=re.S)
    return '\n'.join(line for line in text.splitlines() if not re.search(r'OUTPUT LANGUAGE REQUIREMENT|META INSTRUCTION|Do not translate source names',line,re.I))

def _markdown_blocks(markdown):
    lines=_clean_text(markdown).splitlines();i=0
    while i<len(lines):
        line=lines[i].strip();i+=1
        if not line or re.fullmatch(r'[-*_]{3,}',line): continue
        if line.startswith('|') and line.endswith('|') and i<len(lines) and re.match(r'^\s*\|?\s*:?-{3,}',lines[i]):
            headers=[s.strip() for s in line.strip('|').split('|')];i+=1;rows=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                cells=[s.strip() for s in lines[i].strip().strip('|').split('|')];i+=1
                if len(cells)==len(headers): rows.append(cells)
            if rows: yield ('table',headers,rows)
        elif re.match(r'^#{1,6}\s+',line):
            m=re.match(r'^(#{1,6})\s+(.*)',line);yield ('heading',min(len(m[1]),3),m[2])
        elif re.match(r'^[-*]\s+',line): yield ('bullet',re.sub(r'^[-*]\s+','',line))
        elif re.match(r'^\d+[.)]\s+',line): yield ('number',re.sub(r'^\d+[.)]\s+','',line))
        else: yield ('paragraph',line.removeprefix('> ').strip())

def _inline_doc(p,text):
    # Preserve bold text, render ordinary links as readable text plus URL.
    text=re.sub(r'\[([^\]]+)\]\((https?://[^)]+)\)',r'\1 (\2)',text)
    for i,part in enumerate(re.split(r'\*\*(.*?)\*\*',text)):
        run=p.add_run(part.replace('`',''));run.bold=bool(i%2)

def _inline_pdf(text):
    text=re.sub(r'\[([^\]]+)\]\((https?://[^)]+)\)',r'\1 (\2)',str(text))
    return re.sub(r'\*\*(.*?)\*\*',r'<b>\1</b>',escape(text.replace('`','')))

def _column_weights(headers,rows):
    sizes=[]
    for j,h in enumerate(headers):
        lengths=[len(re.sub(r'\[S\d+\]','',str(r[j]))) for r in rows]
        sizes.append(min(34,max(8,len(h),sorted(lengths)[int((len(lengths)-1)*.7)])))
    return [s/sum(sizes) for s in sizes]

def _set_doc_fonts(doc,language='zh'):
    for name in ('Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','List Bullet','List Number'):
        s=doc.styles[name];s.font.name='Aptos';s.font.color.rgb=RGBColor(0,0,0)
        s._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Noto Sans SC')
        s.paragraph_format.space_after=Pt(7)
    doc.styles['Normal'].font.size=Pt(10.5)
    doc.styles['Normal'].paragraph_format.line_spacing=1.18
    for name,size in [('Title',23),('Heading 1',15),('Heading 2',13),('Heading 3',11)]:
        doc.styles[name].font.size=Pt(size)
        doc.styles[name].paragraph_format.keep_with_next=True
    for style in doc.styles:
        for border in list(style._element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    section=doc.sections[0];section.page_width=Inches(8.27);section.page_height=Inches(11.69)
    section.top_margin=section.bottom_margin=Inches(.7)
    section.left_margin=section.right_margin=Inches(.75)

def _doc_table(doc,headers,rows):
    table=doc.add_table(rows=1,cols=len(headers));table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
    widths=_column_weights(headers,rows)
    for col,w in zip(table.columns,widths): col.width=Inches(6.77*w)
    for idx,values in enumerate([headers]+rows):
        cells=table.rows[0].cells if idx==0 else table.add_row().cells
        pr=table.rows[idx]._tr.get_or_add_trPr()
        if idx==0:
            repeat=OxmlElement('w:tblHeader');pr.append(repeat)
        # Avoid splitting an individual row; allow the table to span pages.
        no_split=OxmlElement('w:cantSplit');pr.append(no_split)
        for j,(cell,value) in enumerate(zip(cells,values)):
            cell.width=Inches(6.77*widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=cell._tc.get_or_add_tcPr();fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'17365D' if idx==0 else ('F0F4F8' if idx%2 else 'FFFFFF'));cp.append(fill)
            borders=OxmlElement('w:tcBorders')
            for side in ('top','left','bottom','right'):
                e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            cp.append(borders)
            margins=OxmlElement('w:tcMar')
            for side in ('top','left','bottom','right'):
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'95');e.set(qn('w:type'),'dxa');margins.append(e)
            cp.append(margins)
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2)
            _inline_doc(p,str(value))
            for run in p.runs:
                run.font.size=Pt(9)
                if idx==0: run.bold=True;run.font.color.rgb=RGBColor(255,255,255)
    doc.add_paragraph().paragraph_format.space_after=Pt(2)

def _add_markdown_to_docx(doc,markdown):
    for block in _markdown_blocks(markdown):
        if block[0]=='heading': doc.add_heading(block[2],level=max(1,block[1]-1))
        elif block[0]=='table': _doc_table(doc,block[1],block[2])
        else:
            p=doc.add_paragraph(style={'bullet':'List Bullet','number':'List Number'}.get(block[0],'Normal'))
            _inline_doc(p,block[1])

def _chart_png(chart):
    labels=chart.get('labels',[]);values=chart.get('values',[])
    if len(labels)<2 or len(labels)!=len(values): return None
    try: values=[float(x) for x in values]
    except (ValueError,TypeError): return None
    alltext=' '.join(labels)+chart.get('title','')+chart.get('unit','')
    font=None
    if re.search(r'[\u3400-\u9fff]',alltext):
        if FONT_PATH.exists():
            font=font_manager.FontProperties(fname=str(FONT_PATH))
        available={f.name for f in font_manager.fontManager.ttflist}
        chosen=next((n for n in ('Noto Sans CJK SC','Noto Sans CJK JP','Microsoft YaHei','SimHei','WenQuanYi Zen Hei') if n in available),None)
        if font is None:
            if not chosen: return None # Never export charts containing missing CJK glyphs.
            font=font_manager.FontProperties(family=chosen)
    fig,ax=plt.subplots(figsize=(7.0,3.8))
    try:
        if chart.get('type')=='line': ax.plot(range(len(labels)),values,marker='o',color='#17365D')
        else: ax.bar(range(len(labels)),values,color='#17365D')
        ax.set_xticks(range(len(labels)));ax.set_xticklabels(labels,fontproperties=font,rotation=20,ha='right')
        ax.set_title(chart.get('title',''),fontproperties=font);ax.set_ylabel(chart.get('unit',''),fontproperties=font)
        ax.grid(axis='y',alpha=.15);fig.tight_layout();bio=BytesIO();fig.savefig(bio,format='png',dpi=150);return bio.getvalue()
    finally: plt.close(fig)

def _pdf_font(language):
    if not FONT_PATH.exists():
        raise FileNotFoundError("Missing data/fonts/NotoSansSC-Regular.ttf; upload the supplied font file.")
    if 'MarketNoto' not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont('MarketNoto',str(FONT_PATH)))
        pdfmetrics.registerFontFamily('MarketNoto',normal='MarketNoto',bold='MarketNoto',italic='MarketNoto',boldItalic='MarketNoto')
    return 'MarketNoto'

def generate_market_docx(result,language='zh'):
    doc=Document();_set_doc_fonts(doc,language);lab=_labels(language)
    doc.add_paragraph(f"{result.get('product','Market Research')}  {result.get('country','')}",style='Title')
    doc.add_paragraph(lab['subtitle'],style='Subtitle')
    date=(result.get('generated_at') or datetime.now().isoformat())[:10]
    doc.add_paragraph(f"{lab['date']}  {date}\n{lab['question']}  {_public_objective(result)}")
    _add_markdown_to_docx(doc,result.get('analysis',''))
    for c in result.get('charts',[]):
        png=_chart_png(c)
        if png:
            doc.add_heading(c.get('title',''),level=2);doc.add_picture(BytesIO(png),width=Inches(6.5))
            doc.add_paragraph(c.get('note','')+' '+', '.join(c.get('source_ids',[])))
        else:
            _doc_table(doc,['Item','Value '+c.get('unit','')],[[str(a),str(b)] for a,b in zip(c.get('labels',[]),c.get('values',[]))])
    used=set(re.findall(r'\[(S\d+)\]',result.get('analysis','')))|{sid for c in result.get('charts',[]) for sid in c.get('source_ids',[])}
    evidence=[(i,s) for i,s in enumerate(result.get('evidence',[]),1) if _source_id(s,i) in used]
    if evidence:
        doc.add_page_break();doc.add_heading(lab['references'],level=1)
        for i,s in evidence:
            p=doc.add_paragraph();p.add_run(f"[{_source_id(s,i)}] {s.get('title','')}").bold=True
            p.add_run('\n'+s.get('link',''))
            if s.get('date'): p.add_run('\n'+s['date'])
    bio=BytesIO();doc.save(bio);return bio.getvalue()

def _pdf_styles(font):
    base=getSampleStyleSheet()
    styles={}
    for name,size,leading in [('body',10,15),('title',22,28),('sub',12,18),('h1',15,21),('h2',13,18),('h3',11,16),('cell',8.6,13),('ref',8.5,12)]:
        styles[name]=ParagraphStyle('Market'+name,fontName=font,fontSize=size,leading=leading,spaceAfter=7,
          spaceBefore=8 if name.startswith('h') else 0,keepWithNext=name.startswith('h'),wordWrap='CJK',splitLongWords=True)
    return styles

def _pdf_table(headers,rows,styles,width=510):
    weights=_column_weights(headers,rows)
    header=ParagraphStyle('MarketHeader',parent=styles['cell'],textColor=colors.white)
    data=[[Paragraph(_inline_pdf(v),header) for v in headers]]+[[Paragraph(_inline_pdf(v),styles['cell']) for v in row] for row in rows]
    table=Table(data,colWidths=[width*w for w in weights],repeatRows=1,hAlign='LEFT')
    table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#17365D')),
      ('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.HexColor('#F0F4F8'),colors.white]),
      ('GRID',(0,0),(-1,-1),.4,colors.HexColor('#D9D9D9')),('VALIGN',(0,0),(-1,-1),'MIDDLE'),
      ('LEFTPADDING',(0,0),(-1,-1),7),('RIGHTPADDING',(0,0),(-1,-1),7),
      ('TOPPADDING',(0,0),(-1,-1),6),('BOTTOMPADDING',(0,0),(-1,-1),6)]))
    return table

def generate_market_pdf(result,language='zh'):
    bio=BytesIO();font=_pdf_font(language);styles=_pdf_styles(font);lab=_labels(language)
    pdf=SimpleDocTemplate(bio,pagesize=A4,rightMargin=42,leftMargin=42,topMargin=42,bottomMargin=42)
    width=A4[0]-84
    story=[Paragraph(_inline_pdf(f"{result.get('product','Market Research')}  {result.get('country','')}"),styles['title']),
      Paragraph(lab['subtitle'],styles['sub']),
      Paragraph(_inline_pdf(lab['date']+'  '+(result.get('generated_at') or datetime.now().isoformat())[:10]),styles['body']),
      Paragraph(_inline_pdf(lab['question']+'  '+_public_objective(result)),styles['body']),Spacer(1,8)]
    for block in _markdown_blocks(result.get('analysis','')):
        if block[0]=='heading': story.append(Paragraph(_inline_pdf(block[2]),styles['h'+str(max(1,block[1]-1))]))
        elif block[0]=='table': story.extend([_pdf_table(block[1],block[2],styles,width),Spacer(1,10)])
        else: story.append(Paragraph(_inline_pdf(('• ' if block[0]=='bullet' else '')+block[1]),styles['body']))
    for c in result.get('charts',[]):
        story.append(Paragraph(_inline_pdf(c.get('title','')),styles['h2']));png=_chart_png(c)
        if png: story.append(RLImage(BytesIO(png),width=width,height=width*3.8/7))
        else: story.append(_pdf_table(['Item','Value '+c.get('unit','')],[[str(a),str(b)] for a,b in zip(c.get('labels',[]),c.get('values',[]))],styles,width))
        story.append(Paragraph(_inline_pdf(c.get('note','')+' '+', '.join(c.get('source_ids',[]))),styles['ref']))
    used=set(re.findall(r'\[(S\d+)\]',result.get('analysis','')))|{sid for c in result.get('charts',[]) for sid in c.get('source_ids',[])}
    evidence=[(i,s) for i,s in enumerate(result.get('evidence',[]),1) if _source_id(s,i) in used]
    if evidence:
        story.extend([PageBreak(),Paragraph(lab['references'],styles['h1'])])
        for i,s in evidence:
            story.append(Paragraph(_inline_pdf(f"[{_source_id(s,i)}] {s.get('title','')}"),styles['ref']))
            url=escape(s.get('link',''));date=escape(s.get('date',''))
            story.append(Paragraph(f'<link href="{url}">{url}</link><br/>{date}',styles['ref']))
    def page_number(canvas,doc):
        canvas.saveState();canvas.setFont('Helvetica',8);canvas.setFillColor(colors.HexColor('#667085'))
        canvas.drawRightString(A4[0]-42,25,str(doc.page));canvas.restoreState()
    pdf.build(story,onFirstPage=page_number,onLaterPages=page_number)
    return bio.getvalue()
